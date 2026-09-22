#!/usr/bin/env python3
"""Build reference.docx for this doctype: pandoc's default, patched to the plain
BFS-Aarau look from Jonas' example (Arial, black headings, header with title/author,
footer with school + page number). Rerun after changing anything here.
    /usr/bin/python3 doctypes/lernbericht-bfs-aarau/make-reference.py
"""
import re
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).parent
base = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
tmp = HERE / "_base.docx"
tmp.write_bytes(base)


def patch_styles(xml: bytes) -> bytes:
    arial = b'<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/>'
    xml = re.sub(rb"<w:rFonts\b[^>]*/>", arial, xml, flags=re.S)
    xml = re.sub(rb'<w:color w:val="[0-9A-F]{6}"\s+w:themeColor="accent1"[^>]*/>', b'<w:color w:val="000000"/>', xml, flags=re.S)
    return xml


with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(HERE / "_styled.docx", "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/styles.xml":
            data = patch_styles(data)
        zout.writestr(item, data)

doc = Document(str(HERE / "_styled.docx"))
section = doc.sections[0]
section.header.is_linked_to_previous = False
section.footer.is_linked_to_previous = False

header_p = section.header.paragraphs[0]
header_p.text = "Lernbericht · Spritztechnik · Lia Brunner"
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in header_p.runs:
    run.font.size = doc.styles["Normal"].font.size
    run.font.italic = True

footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.text = ""
run = footer_p.add_run("BFS Aarau · Lernbericht · L. Brunner · Seite ")
run.font.italic = True

fld_begin = OxmlElement("w:fldChar")
fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = "PAGE"
fld_sep = OxmlElement("w:fldChar")
fld_sep.set(qn("w:fldCharType"), "separate")
fld_end = OxmlElement("w:fldChar")
fld_end.set(qn("w:fldCharType"), "end")

page_run = footer_p.add_run()
page_run.font.italic = True
page_run._r.append(fld_begin)
page_run2 = footer_p.add_run()
page_run2._r.append(instr)
page_run3 = footer_p.add_run()
page_run3._r.append(fld_sep)
page_run4 = footer_p.add_run()
page_run4._r.append(fld_end)

doc.save(str(HERE / "reference.docx"))
tmp.unlink()
(HERE / "_styled.docx").unlink()
print("✔ reference.docx")
